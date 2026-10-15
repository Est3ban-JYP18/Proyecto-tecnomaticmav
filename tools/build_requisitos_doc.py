from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r"E:\Proyecto-tecnomaticmav-limpio\Proyecto-tecnomaticmav\Documentacion_Requisitos_Tecnomatic_MAV.docx")
BLUE, LIGHT_BLUE, GRAY = "123B66", "EAF2F8", "D9D9D9"

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd"); el.set(qn("w:fill"), fill); pr.append(el)

def border(cell):
    pr = cell._tc.get_or_add_tcPr(); borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:color"), GRAY); borders.append(el)
    pr.append(borders)

def margins(cell):
    pr = cell._tc.get_or_add_tcPr(); mar = OxmlElement("w:tcMar")
    for key in ("top", "start", "bottom", "end"):
        el = OxmlElement(f"w:{key}"); el.set(qn("w:w"), "100"); el.set(qn("w:type"), "dxa"); mar.append(el)
    pr.append(mar)

def font(run, size=10, white=False):
    run.font.name = "Aptos"; run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos"); run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size); run.font.color.rgb = RGBColor(255,255,255) if white else RGBColor(0,0,0)

def add_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, label in enumerate(headers):
        c = t.rows[0].cells[i]; c.width = Inches(widths[i]); shade(c, BLUE); border(c); margins(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(label); r.bold = True; font(r, 8.5, True)
    trpr = t.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); trpr.append(repeat)
    for n, values in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(values):
            c = cells[i]; c.width = Inches(widths[i]); shade(c, "FFFFFF" if n % 2 == 0 else LIGHT_BLUE); border(c); margins(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05
            if i == 0: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value)); r.bold = i == 0; font(r, 8.4)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def para(doc, value, bold_prefix=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.12
    if bold_prefix and value.startswith(bold_prefix):
        r = p.add_run(bold_prefix); r.bold = True; font(r); r = p.add_run(value[len(bold_prefix):]); font(r)
    else:
        font(p.add_run(value))
    return p

def heading(doc, value, level=1):
    p = doc.add_paragraph(style=f"Heading {level}"); p.paragraph_format.space_before = Pt(12 if level == 1 else 8); p.paragraph_format.space_after = Pt(5); p.paragraph_format.keep_with_next = True
    font(p.add_run(value), 15 if level == 1 else 12)
    return p

def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3); font(p.add_run(value)); return p

def page_field(p):
    r = p.add_run(); begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instr, end])

doc = Document()
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.68)
sec.left_margin, sec.right_margin = Inches(.78), Inches(.78)
for name, size in (("Normal",10), ("Title",24), ("Heading 1",15), ("Heading 2",12)):
    st = doc.styles[name]; st.font.name = "Aptos"; st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos"); st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos"); st.font.size = Pt(size); st.font.color.rgb = RGBColor(0,0,0)

hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; r = hp.add_run("TECNOMATIC MAV | Especificación de requisitos"); font(r,8); r.font.color.rgb = RGBColor(90,90,90)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = fp.add_run("Documento de requisitos | Página "); font(r,8); r.font.color.rgb = RGBColor(90,90,90); page_field(fp)

p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(80); p.paragraph_format.space_after = Pt(10); p.add_run("Especificación de requisitos del sistema")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(28); r = p.add_run("Tecnomatic MAV"); r.bold = True; font(r,17)
add_table(doc, ["Elemento","Descripción"], [
 ("Proyecto","Plataforma web de comercio electrónico y administración de dotación y elementos de protección industrial."),
 ("Propósito","Definir el comportamiento esperado, los límites de calidad y los criterios para validar el sistema."),
 ("Usuarios","Visitantes, clientes, administradores y contadores o gestores de inventario."),
 ("Persistencia","Supabase PostgreSQL para catálogo, usuarios, inventario, pedidos, devoluciones, cupones y cotizaciones."),
 ("Versión","1.0")], [1.3,5.8])
para(doc, "Este documento establece los requisitos funcionales y no funcionales de Tecnomatic MAV. Su finalidad es servir como referencia para desarrollo, pruebas, entrega académica y evolución del sistema. Los requisitos se derivan de los módulos implementados en la aplicación: catálogo, carrito, autenticación, administración, inventario, pedidos, cotizaciones, devoluciones y configuración de Supabase.")
para(doc, "Conclusión principal: el sistema debe permitir gestionar de forma centralizada un catálogo de productos de seguridad industrial y sus procesos comerciales, conservando la información en Supabase y aplicando permisos acordes con el rol del usuario.")
doc.add_page_break()

heading(doc,"1 Objetivo y alcance")
para(doc,"Tecnomatic MAV es una aplicación web para la consulta y venta de elementos de protección personal, dotación industrial y productos relacionados. El sistema combina una interfaz de cliente con módulos administrativos para mantener productos, usuarios, inventario, pedidos, cotizaciones y devoluciones.")
heading(doc,"1.1 Alcance incluido",2)
for x in ["Consulta de catálogo público, filtros, fichas técnicas, imágenes, variantes y disponibilidad.","Registro, inicio y cierre de sesión con perfiles de administrador, contador y cliente.","Carrito de compras, aplicación de cupones, creación de pedidos y seguimiento de entregas.","Administración de productos, categorías, stock, movimientos de inventario y alertas de reposición.","Gestión de usuarios, facturas, cotizaciones, reseñas y solicitudes de devolución.","Persistencia de la información operativa en Supabase PostgreSQL."]: bullet(doc,x)
heading(doc,"1.2 Fuera de alcance",2)
for x in ["Pasarela de pago real, conciliación bancaria o generación fiscal de facturas electrónicas.","Integración automática con transportadoras, correo transaccional o proveedores externos.","Aplicación móvil nativa y operación sin conexión como modo principal."]: bullet(doc,x)
heading(doc,"2 Actores y permisos")
add_table(doc,["Actor","Responsabilidades dentro del sistema"],[
 ("Visitante","Consulta el catálogo público, filtra productos y puede iniciar un registro o solicitar una cotización."),
 ("Cliente","Administra su carrito, crea pedidos, consulta sus pedidos, solicita devoluciones y publica reseñas."),
 ("Contador","Consulta indicadores y apoya procesos de inventario, pedidos y control operativo según los permisos asignados."),
 ("Administrador","Gestiona usuarios, productos, catálogo técnico, inventario, facturas, cotizaciones, devoluciones y estados operativos.")],[1.35,5.75])
heading(doc,"3 Requisitos funcionales")
para(doc,"Los siguientes requisitos describen funciones observables del sistema. El identificador permite relacionarlos con pruebas y entregables.")
add_table(doc,["ID","Módulo","Requisito"],[
 ("RF 01","Autenticación","El sistema debe permitir registro, inicio de sesión, cierre de sesión y recuperación de contraseña. Debe identificar el rol del usuario al iniciar sesión."),
 ("RF 02","Gestión de usuarios","El administrador debe poder consultar, crear, actualizar y eliminar usuarios, asignando el rol correspondiente."),
 ("RF 03","Catálogo público","El sistema debe mostrar productos activos con nombre, precio, imagen, categoría, stock y atributos relevantes."),
 ("RF 04","Búsqueda y filtros","El usuario debe poder filtrar productos por categoría, rango de precio, disponibilidad, marca, material y nivel de protección."),
 ("RF 05","Ficha técnica","El sistema debe mostrar atributos, galería de imágenes, variantes de color y talla, reseñas y productos relacionados."),
 ("RF 06","Carrito","El cliente debe agregar productos, seleccionar talla y color, modificar cantidades, eliminar artículos y vaciar el carrito respetando el stock disponible."),
 ("RF 07","Cupones","El sistema debe validar códigos activos, verificar compra mínima y calcular el descuento porcentual o fijo antes de crear el pedido.")],[.65,1.4,5.05])
doc.add_page_break()

heading(doc,"3 Requisitos funcionales continuación")
add_table(doc,["ID","Módulo","Requisito"],[
 ("RF 08","Pedidos","El cliente debe poder confirmar un pedido. El sistema debe registrar factura, detalle, entrega, descuento de inventario y movimiento de salida."),
 ("RF 09","Seguimiento","El cliente debe consultar el estado, fecha estimada, observaciones y referencia de seguimiento de cada pedido."),
 ("RF 10","Productos","El administrador debe crear, editar, desactivar y enviar productos a papelera, incluyendo categoría, precio, imagen y estado."),
 ("RF 11","Catálogo técnico","El administrador debe administrar atributos técnicos, galería de imágenes y variantes de producto."),
 ("RF 12","Inventario","El sistema debe visualizar stock actual, alertas de stock bajo, movimientos por producto y ajustes de entrada o salida."),
 ("RF 13","Ajustes masivos","El administrador debe poder actualizar precios por porcentaje o stock por variación para todos los productos de una categoría."),
 ("RF 14","Cotizaciones","El sistema debe registrar solicitudes de cotización, mostrar sus productos solicitados y permitir actualizar su estado."),
 ("RF 15","Devoluciones","El cliente debe crear una solicitud de devolución; el administrador debe aprobarla, rechazarla, registrar comentarios y actualizar el seguimiento."),
 ("RF 16","Persistencia","Las acciones operativas deben guardarse en Supabase y reflejarse en las consultas posteriores del módulo correspondiente.")],[.65,1.4,5.05])
heading(doc,"4 Reglas de negocio")
add_table(doc,["ID","Regla"],[
 ("RN 01","Todo producto debe tener nombre, precio no negativo y una categoría asociada."),
 ("RN 02","Un producto con cantidad de stock igual a cero debe quedar en estado Agotado."),
 ("RN 03","No se puede agregar al carrito una cantidad mayor a la disponible en inventario."),
 ("RN 04","Un cupón solo se aplica cuando está activo y se cumple su compra mínima."),
 ("RN 05","Un pedido requiere un usuario cliente y al menos un detalle de producto."),
 ("RN 06","Las variantes no pueden repetir la misma combinación de producto, color y talla."),
 ("RN 07","Las solicitudes de devolución deben conservar la relación con factura, producto, cliente y motivo.")],[.75,6.35])
doc.add_page_break()

heading(doc,"5 Requisitos no funcionales")
para(doc,"Estos requisitos establecen atributos de calidad y condiciones técnicas que el sistema debe cumplir, además de sus funciones.")
add_table(doc,["ID","Atributo","Requisito verificable"],[
 ("RNF 01","Seguridad","Las claves de Supabase expuestas al navegador deben limitarse a la Anon Key. La Service Role Key no puede incluirse en el cliente."),
 ("RNF 02","Control de acceso","Supabase debe aplicar políticas RLS. En producción, clientes solo pueden consultar y modificar información propia; los permisos administrativos deben requerir un rol autorizado."),
 ("RNF 03","Privacidad","Las contraseñas no deben ser devueltas al navegador ni conservarse en texto plano. La versión productiva debe utilizar Supabase Auth o un mecanismo equivalente con hash seguro."),
 ("RNF 04","Rendimiento","Las vistas principales del catálogo y del inventario deben cargar en menos de 3 segundos bajo una conexión de banda ancha y un conjunto de hasta 1.000 productos."),
 ("RNF 05","Usabilidad","Las acciones críticas deben informar éxito, error o validación mediante mensajes visibles y comprensibles en español."),
 ("RNF 06","Compatibilidad","La interfaz debe operar en versiones recientes de Chrome, Edge, Firefox y navegadores móviles modernos, con diseño adaptable."),
 ("RNF 07","Disponibilidad","Cuando Supabase no esté disponible, la aplicación debe informar el problema y no presentar como persistidos datos que solo quedaron en almacenamiento local."),
 ("RNF 08","Integridad","Las claves foráneas y restricciones de la base de datos deben impedir detalles de pedido sin factura, stock sin producto o categorías duplicadas."),
 ("RNF 09","Mantenibilidad","Las consultas a Supabase deben concentrarse en servicios reutilizables y conservar nombres de tablas y columnas consistentes con el esquema implementado."),
 ("RNF 10","Trazabilidad","Los cambios de inventario deben generar un movimiento con tipo, cantidad, fecha y observación; los cambios administrativos relevantes deben poder auditarse."),
 ("RNF 11","Respaldo","La base de datos debe contar con copias de seguridad y un procedimiento de recuperación definido desde el panel de Supabase."),
 ("RNF 12","Escalabilidad","Las tablas usadas para búsquedas y políticas deben disponer de índices en claves primarias, correos, claves foráneas y campos de filtrado frecuente.")],[.65,1.35,5.1])
doc.add_page_break()

heading(doc,"6 Datos y persistencia")
para(doc,"La solución utiliza Supabase PostgreSQL como fuente de persistencia. El esquema incluye entidades para roles, usuarios, categorías, productos, stock, movimientos de inventario, atributos, imágenes, variantes, facturas, detalles, entregas, devoluciones, cupones, alertas y cotizaciones.")
add_table(doc,["Área","Entidades principales","Información preservada"],[
 ("Catálogo","Categoria_productos, Productos, Producto_Atributos, Producto_Imagenes, Producto_Variantes","Identidad, precio, imágenes, especificaciones, categoría y variantes."),
 ("Inventario","Stock, Movimientos_Inventario, Alertas_Reposicion","Cantidad actual, ajustes, entradas, salidas y alertas."),
 ("Comercial","Facturas, Detalle_Facturas, Entregas, Cupones, Cotizaciones","Pedido, artículos, estados, descuentos, entrega y solicitudes comerciales."),
 ("Servicio","Devoluciones, Resenas_Productos, Usuarios, Roles","Solicitudes posteriores a la venta, evaluación, identidad y autorización.")],[1.05,2.75,3.3])
heading(doc,"7 Criterios de aceptación prioritarios")
add_table(doc,["ID","Resultado esperado"],[
 ("CA 01","Al iniciar sesión con una cuenta válida, el sistema muestra las opciones correspondientes a su rol."),
 ("CA 02","Al crear un producto, este se visualiza en el catálogo y se registra su stock inicial."),
 ("CA 03","Al modificar el stock, el valor se actualiza y aparece un movimiento de inventario."),
 ("CA 04","Al confirmar un pedido, se crean factura, detalle y entrega, y el stock disminuye de forma coherente."),
 ("CA 05","Al aplicar un cupón válido, el total final refleja el descuento y se rechaza el cupón si no cumple el mínimo."),
 ("CA 06","Al enviar una devolución, el cliente puede verla; al aprobarla o rechazarla, el estado queda actualizado."),
 ("CA 07","Una operación denegada por permisos muestra un mensaje comprensible y no altera datos."),
 ("CA 08","Tras recargar la página, los datos registrados en Supabase permanecen disponibles.")],[.75,6.35])
heading(doc,"8 Recomendaciones para la entrega")
para(doc,"Para la demostración actual, los permisos de desarrollo permiten validar la persistencia de cada módulo. Antes de publicar el sistema se debe reemplazar esa configuración por políticas RLS basadas en Supabase Auth, retirar contraseñas de la tabla pública y limitar los permisos por rol. Esta transición protege clientes, pedidos e inventario sin cambiar el alcance funcional documentado.")

doc.core_properties.title = "Especificación de requisitos del sistema Tecnomatic MAV"
doc.core_properties.subject = "Requisitos funcionales y no funcionales"
doc.core_properties.author = "Tecnomatic MAV"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)

